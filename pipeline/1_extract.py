import os as _os
_ROOT=_os.path.dirname(_os.path.dirname(_os.path.abspath(__file__)))
_DATA=_os.path.join(_ROOT,"data")
_SRC=_os.path.join(_ROOT,"source")
import openpyxl, json, datetime
f = _os.path.join(_SRC,"AFable_EasyHardMoney_NDX十年每日分段_2026R4.5.3_2.xlsx")
wb = openpyxl.load_workbook(f, data_only=True, read_only=True)
ws = wb['02A_每日分段']
# columns: B=date(2), C=year(3), E=NDX(5), Y=zone(25), AX=score(50), S=dailychg(19), AB=policycode(28)
def cidx(letter):
    return openpyxl.utils.column_index_from_string(letter)
cols = {'date':2,'year':3,'ndx':5,'zone':25,'score':50,'chg':19,'pol':28}
data=[]
for row in ws.iter_rows(min_row=5, max_row=ws.max_row, values_only=True):
    d = row[cols['date']-1]
    if d is None: continue
    if isinstance(d, datetime.datetime): d=d.date()
    z = row[cols['zone']-1]
    ndx = row[cols['ndx']-1]
    score = row[cols['score']-1]
    if z is None: continue
    data.append({'date':d.isoformat(),'year':row[cols['year']-1],'ndx':ndx,'zone':z,'score':score,'pol':row[cols['pol']-1]})
print("total daily rows:", len(data))
print("first:", data[0])
print("last:", data[-1])
# zone distribution
from collections import Counter
print("zone counts:", Counter(x['zone'] for x in data))
# save
with open(_os.path.join(_DATA,"daily.json"),'w') as fp:
    json.dump(data, fp)
print("saved daily.json")
