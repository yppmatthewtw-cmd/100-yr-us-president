import os as _os
_ROOT=_os.path.dirname(_os.path.dirname(_os.path.abspath(__file__)))
_DATA=_os.path.join(_ROOT,"data")
_SRC=_os.path.join(_ROOT,"source")
import openpyxl, json, datetime
# --- R1 historical table (1925-2026) ---
f1=_os.path.join(_SRC,"R1_百年逐屆EasyHard資金市場週期.xlsx")
wb1=openpyxl.load_workbook(f1,data_only=True)
ws=wb1['1.百年逐屆EasyHard']
hist=[]
for r in range(4,30):
    a=ws.cell(r,1).value
    if not a: continue
    hist.append({'pres':str(a).replace('\n',' '),'easy':ws.cell(r,2).value,'hard':ws.cell(r,3).value,'events':ws.cell(r,4).value})
# stats sheet
ws2=wb1['2.統計規律']
cycstats=[]
for r in range(4,8):
    if ws2.cell(r,1).value:
        cycstats.append({'year':ws2.cell(r,1).value,'ret':ws2.cell(r,2).value,'upratio':ws2.cell(r,3).value,'feat':ws2.cell(r,4).value})

# --- L1 policy steps (07 sheet) ---
f2=_os.path.join(_SRC,"AFable_EasyHardMoney_NDX十年每日分段_2026R4.5.3_2.xlsx")
wb2=openpyxl.load_workbook(f2,data_only=True)
ws3=wb2['07_L1政策步階']
pol=[]
for r in range(4,15):
    d=ws3.cell(r,1).value
    if d is None: continue
    if isinstance(d,datetime.datetime): d=d.date().isoformat()
    pol.append({'date':d,'state':ws3.cell(r,2).value,'code':ws3.cell(r,3).value,'basis':ws3.cell(r,5).value})

# --- indicator library (01 sheet) categories ---
ws4=wb2['01_指標庫']
indics=[]
for r in range(4,22):
    if ws4.cell(r,1).value:
        indics.append({'id':ws4.cell(r,1).value,'cat':ws4.cell(r,2).value,'name':ws4.cell(r,3).value,'weight':ws4.cell(r,8).value,'desc':ws4.cell(r,9).value})

pkg={'hist':hist,'cycstats':cycstats,'l1policy':pol,'indicators':indics}
pkg['segments']=json.load(open(_os.path.join(_DATA,"segments.json")))
pkg['cycleyears']=json.load(open(_os.path.join(_DATA,"cycleyears.json")))
json.dump(pkg, open(_os.path.join(_DATA,"package.json"),'w'), ensure_ascii=False, indent=1)
print("HIST terms:", len(hist))
print("L1 policy steps:", len(pol))
print("indicators:", len(indics))
print("cyclestats years:", len(cycstats))
print("\nL1 POLICY STEPS:")
for p in pol: print(" ", p['date'], p['code'], str(p['state'])[:55])
print("\nsaved package.json (size KB):", round(len(open(_os.path.join(_DATA,"package.json")).read())/1024,1))
